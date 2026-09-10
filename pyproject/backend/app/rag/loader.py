import csv
import io
from pathlib import Path

from pypdf import PdfReader

from app.config import MAX_CSV_ROWS
from app.rag.textutils import is_heading_text


class Document:
    def __init__(self, page_content: str, metadata: dict | None = None):
        self.page_content = page_content
        self.metadata = metadata or {}


def _flush_section(sections: list[dict], heading: str, parts: list[str], source: str, page: int) -> None:
    content = "\n".join(p for p in parts if p.strip()).strip()
    if not content:
        return
    sections.append({"heading": heading, "content": content, "source": source, "page": page})


def _sections_to_documents(sections: list[dict]) -> list[Document]:
    docs: list[Document] = []
    for item in sections:
        heading = item["heading"]
        body = item["content"]
        page_content = body if body.startswith(heading) else f"{heading}\n{body}"
        docs.append(
            Document(
                page_content,
                {
                    "source": item["source"],
                    "page": item["page"],
                    "heading": heading,
                    "kind": item.get("kind", "section"),
                    "record_id": item.get("record_id", heading),
                },
            )
        )
    return docs


def load_txt(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    sections: list[dict] = []
    heading = path.stem.replace("_", " ")
    parts: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if is_heading_text(line) and parts:
            _flush_section(sections, heading, parts, path.name, 1)
            heading = line
            parts = [line]
        else:
            parts.append(line)
    _flush_section(sections, heading, parts, path.name, 1)
    if sections:
        return _sections_to_documents(sections)
    return [Document(text, {"source": path.name, "page": 1, "heading": path.stem})]


def load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(str(path))
    sections: list[dict] = []
    heading = path.stem.replace("_", " ")
    parts: list[str] = []
    page_for_heading = 1
    for i, page in enumerate(reader.pages, start=1):
        content = page.extract_text() or ""
        for raw in content.splitlines():
            line = raw.strip()
            if not line:
                continue
            if is_heading_text(line) and parts:
                _flush_section(sections, heading, parts, path.name, page_for_heading)
                heading = line
                parts = [line]
                page_for_heading = i
            else:
                parts.append(line)
    _flush_section(sections, heading, parts, path.name, page_for_heading)
    if sections:
        return _sections_to_documents(sections)
    docs: list[Document] = []
    for i, page in enumerate(reader.pages, start=1):
        content = page.extract_text() or ""
        if content.strip():
            docs.append(Document(content, {"source": path.name, "page": i, "heading": path.stem}))
    return docs


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        cells = [c for c in cells if c]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _iter_docx_blocks(doc):
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(str(path))
    sections: list[dict] = []
    heading = path.stem.replace("_", " ")
    parts: list[str] = []

    for block in _iter_docx_blocks(doc):
        if isinstance(block, Table):
            rendered = _table_text(block)
            if rendered:
                parts.append(rendered)
            continue
        if not isinstance(block, Paragraph):
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name or "").lower() if block.style else ""
        heading_like = "heading" in style or is_heading_text(text)
        if heading_like:
            if parts:
                _flush_section(sections, heading, parts, path.name, 1)
            heading = text.rstrip(":")
            parts = [heading]
        else:
            parts.append(text)

    _flush_section(sections, heading, parts, path.name, 1)
    if sections:
        return _sections_to_documents(sections)
    return [Document("\n".join(parts), {"source": path.name, "page": 1, "heading": path.stem})]


def _pretty_field(name: str) -> str:
    return name.replace("_", " ").strip().capitalize()


def _read_csv_rows(text: str) -> tuple[list[str], list[list[str]]]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        return [], []
    header = [cell.strip() or f"col{i+1}" for i, cell in enumerate(rows[0])]
    return header, rows[1:]


def load_csv(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    header, rows = _read_csv_rows(text)
    if not rows:
        return [Document(text, {"source": path.name, "page": 1, "heading": path.stem, "kind": "section"})]
    truncated = len(rows) > MAX_CSV_ROWS
    rows = rows[:MAX_CSV_ROWS]
    docs: list[Document] = []
    for index, row in enumerate(rows, start=1):
        values = list(row) + [""] * (len(header) - len(row))
        data = {header[i]: values[i].strip() for i in range(len(header))}
        record_id = next((data[key] for key in header if data.get(key)), f"linha {index}")
        lines = [f"{_pretty_field(key)}: {data[key]}" for key in header if data.get(key)]
        content = "\n".join(lines)
        docs.append(
            Document(
                content,
                {
                    "source": path.name,
                    "page": 1,
                    "heading": record_id,
                    "kind": "row",
                    "record_id": record_id,
                    "row": index,
                    "truncated": truncated,
                },
            )
        )
    return docs


LOADERS = {
    ".txt": load_txt,
    ".md": load_txt,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".csv": load_csv,
}


def load_document(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    loader = LOADERS.get(suffix)
    if loader is None:
        raise ValueError(f"Formato nao suportado: {suffix}")
    return loader(path)


def load_directory(directory: Path) -> list[Document]:
    documents: list[Document] = []
    if not directory.exists():
        return documents
    for path in sorted(directory.iterdir()):
        if path.is_file() and path.suffix.lower() in LOADERS:
            documents.extend(load_document(path))
    return documents
